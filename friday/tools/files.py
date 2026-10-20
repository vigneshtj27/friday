"""Stored files (statements, Drive copies) and small table computations, all in plain Python.

table_compute is how friday totals or filters real data: it reads the cells itself (a stored
CSV/XLSX, or a sheet tab) and computes in code, so a number reported to the owner came from the
data, not from the model's arithmetic. No shell, no arbitrary code.
"""

from __future__ import annotations

import csv
import io
import re
from collections import defaultdict
from datetime import date
from decimal import Decimal
from typing import Any

from friday.capabilities.transactions import parse_amount, parse_date, pdf_text
from friday.sheets import find_header, spreadsheet_id
from friday.tools.base import Call, Tool, ToolError

OPS = ("eq", "ne", "contains", "not_contains", "gt", "ge", "lt", "le", "between", "month", "nonempty")


def _rows_from_file(ctx: Any, file_id: str) -> tuple[str, list[list[str]] | None, str | None]:
    meta = ctx.files.meta(file_id)
    if meta is None:
        raise ToolError("no such stored file")
    data = ctx.files.read(file_id)
    name = meta.name.lower()
    if data[:5] == b"%PDF-" or name.endswith(".pdf"):
        return meta.name, None, pdf_text(data)
    if data[:2] == b"PK" or name.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = [["" if v is None else (v.isoformat()[:10] if isinstance(v, date) else str(v)) for v in r]
                for r in wb.active.iter_rows(values_only=True)]  # fmt: skip
        return meta.name, rows, None
    text = data.decode("utf-8-sig", "replace")
    if name.endswith((".csv", ".tsv")) or "," in text.split("\n", 1)[0]:
        delim = "\t" if name.endswith(".tsv") else ","
        return meta.name, [r for r in csv.reader(io.StringIO(text), delimiter=delim)], None
    return meta.name, None, text


def _date(v: str) -> date | None:
    return parse_date(v, dayfirst=True) if v and re.search(r"\d", v) else None


def _num(v: str) -> Decimal | None:
    a = parse_amount(v) if v and re.search(r"\d", v) else None
    if a is None:
        return None
    return -a[0] if v.strip().startswith(("-", "(", "−")) else a[0]


def _match(cell: str, op: str, value: Any) -> bool:
    c = cell.strip()
    if op == "nonempty":
        return bool(c)
    if op == "contains":
        return str(value).lower() in c.lower()
    if op == "not_contains":
        return str(value).lower() not in c.lower()
    if op == "month":
        d = _date(c)
        return d is not None and d.strftime("%Y-%m") == str(value)
    if op == "between":
        lo, hi = value if isinstance(value, list) and len(value) == 2 else (None, None)
        d, dlo, dhi = _date(c), _date(str(lo)), _date(str(hi))
        if d and dlo and dhi:
            return dlo <= d <= dhi
        n, nlo, nhi = _num(c), _num(str(lo)), _num(str(hi))
        return n is not None and nlo is not None and nhi is not None and nlo <= n <= nhi
    if op in ("eq", "ne"):
        same = c.lower() == str(value).strip().lower()
        n, v = _num(c), _num(str(value))
        same = same or (n is not None and v is not None and n == v)
        return same if op == "eq" else not same
    d, dv = _date(c), _date(str(value))
    a, b = (d, dv) if d and dv else (_num(c), _num(str(value)))
    if a is None or b is None:
        return False
    return {"gt": a > b, "ge": a >= b, "lt": a < b, "le": a <= b}[op]


def compute(header: list[str], rows: list[list[str]], *, filters: list[dict[str, Any]], sum_columns: list[str],
            group_by: str | None, sample: int) -> dict[str, Any]:  # fmt: skip
    idx = {h.strip().lower(): i for i, h in enumerate(header) if str(h).strip()}

    def col(name: str) -> int:
        k = name.strip().lower()
        if k not in idx:
            raise ToolError(f"no column {name!r}; columns are {[h for h in header if str(h).strip()]}")
        return idx[k]

    for f in filters:
        if f.get("op") not in OPS:
            raise ToolError(f"filter op must be one of {OPS}")
    cells = lambda r, i: str(r[i]) if i < len(r) else ""  # noqa: E731
    kept = [r for r in rows if any(str(c).strip() for c in r)
            and all(_match(cells(r, col(f["column"])), f["op"], f.get("value")) for f in filters)]  # fmt: skip
    out: dict[str, Any] = {"columns": header, "rows_considered": len([r for r in rows if any(str(c).strip() for c in r)]),
                           "rows_matched": len(kept)}  # fmt: skip
    for s in sum_columns:
        i = col(s)
        vals = [_num(cells(r, i)) for r in kept]
        out.setdefault("sums", {})[s] = str(sum((v for v in vals if v is not None), Decimal(0)))
        skipped = sum(1 for v in vals if v is None)
        if skipped:
            out.setdefault("non_numeric_cells", {})[s] = skipped
    if group_by:
        g = col(group_by)
        groups: dict[str, dict[str, Any]] = defaultdict(lambda: {"rows": 0})
        for r in kept:
            e = groups[cells(r, g).strip() or "(blank)"]
            e["rows"] += 1
            for s in sum_columns:
                v = _num(cells(r, col(s)))
                if v is not None:
                    e[s] = str(Decimal(e.get(s, "0")) + v)
        out["groups"] = dict(sorted(groups.items(), key=lambda kv: -kv[1]["rows"])[:50])
    out["sample_rows"] = kept[:sample]
    return out


class FilesList(Tool):
    name = "files_list"
    description = "Files in friday's encrypted store (downloaded statements, Drive copies)."
    schema = {"type": "object", "properties": {"account": {"type": "string"}, "kind": {"type": "string"}}}

    async def run(self, args: dict[str, Any], call: Call) -> Any:
        files = call.ctx.files.list(profile_id=args.get("account"), kind=args.get("kind"), limit=50)
        return {"files": [{"file_id": f.id, "name": f.name, "mime": f.mime, "size": f.size, "kind": f.kind,
                           "account": f.profile_id, "created_at": f.created_at} for f in files]}  # fmt: skip


class FilesRead(Tool):
    name = "files_read"
    description = "Read a stored file: CSV/XLSX as rows (first `limit`), PDF or text as text."
    schema = {"type": "object", "properties": {"file_id": {"type": "string"}, "limit": {"type": "integer"}}, "required": ["file_id"]}
    untrusted_output = True

    async def run(self, args: dict[str, Any], call: Call) -> Any:
        name, rows, text = _rows_from_file(call.ctx, args["file_id"])
        if rows is not None:
            return {"name": name, "row_count": len(rows), "rows": rows[: int(args.get("limit") or 50)]}
        return {"name": name, "text": (text or "")[:30000], "truncated": len(text or "") > 30000}


class TableCompute(Tool):
    name = "table_compute"
    description = (
        "Filter, total and group a table computed in code from real cells: a stored CSV/XLSX (file_id) or a sheet tab "
        "(spreadsheet + tab). filters: [{column, op, value}] with op eq|ne|contains|not_contains|gt|ge|lt|le|between "
        "(value [lo, hi])|month (value YYYY-MM)|nonempty. sum: column names to total. group_by: a column."
    )
    schema = {"type": "object", "properties": {
        "file_id": {"type": "string"}, "spreadsheet": {"type": "string"}, "tab": {"type": "string"},
        "filters": {"type": "array", "items": {"type": "object", "properties": {
            "column": {"type": "string"}, "op": {"type": "string", "enum": list(OPS)},
            "value": {"type": ["string", "number", "array", "null"]}}, "required": ["column", "op"]}},
        "sum": {"type": "array", "items": {"type": "string"}}, "group_by": {"type": "string"}, "sample": {"type": "integer"},
    }}  # fmt: skip
    untrusted_output = True

    async def run(self, args: dict[str, Any], call: Call) -> Any:
        if args.get("spreadsheet"):
            if not args.get("tab"):
                raise ToolError("which tab?")
            values = await call.ctx.sheets.read(spreadsheet_id(args["spreadsheet"]), args["tab"])
            source = f"sheet tab {args['tab']!r}"
        elif args.get("file_id"):
            name, values, _ = _rows_from_file(call.ctx, args["file_id"])
            if values is None:
                raise ToolError("that file is text, not a table")
            source = f"file {name!r}"
        else:
            raise ToolError("give file_id, or spreadsheet and tab")
        if not values:
            raise ToolError("the table is empty")
        try:
            h = find_header(values)
        except Exception:
            h = 0
        header = [str(c) for c in values[h]]
        out = compute(header, [[str(c) for c in r] for r in values[h + 1 :]], filters=args.get("filters") or [],
                      sum_columns=args.get("sum") or [], group_by=args.get("group_by"), sample=int(args.get("sample") or 5))  # fmt: skip
        out["source"] = source
        out["computed_in_code"] = True
        out["verified"] = [f"computed in code from {out['rows_matched']} matching row(s) of {source}"]
        return out
